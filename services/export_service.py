"""Generates Excel, CSV (zipped), and PDF reports from report_service data.
Everything is produced in-memory (BytesIO) — nothing is written permanently
to disk, and filenames are always application-controlled.
"""
import csv
import io
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from utils.date_utils import format_date_for_display

TASK_HEADERS = ["Date", "Time", "Title", "Priority", "Status", "Due Date", "Notes"]
MEETING_HEADERS = ["Date", "Time", "Title", "My Points", "Meeting Points", "Decisions", "Notes"]


def _task_row(t):
    return [
        format_date_for_display(t["task_date"]),
        t["task_time"] or "",
        t["title"],
        t["priority"],
        t["status"].replace("_", " ").title(),
        format_date_for_display(t["due_date"]) if t["due_date"] else "",
        t["notes"] or "",
    ]


def _meeting_row(m):
    return [
        format_date_for_display(m["meeting_date"]),
        m["meeting_time"] or "",
        m["title"],
        m["my_points"] or "",
        m["meeting_points"] or "",
        m["decisions"] or "",
        m["notes"] or "",
    ]


# ---------------------------------------------------------------- Excel ----

def generate_excel(report_data):
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws["A1"] = "Daily Task & Meeting Tracker - Report"
    summary_ws["A1"].font = Font(bold=True, size=14)
    summary_ws["A2"] = f"Report Period: {report_data['range_label']}"
    summary_ws["A3"] = f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}"

    summary_rows = [
        ("Total Tasks", report_data["summary"]["total_tasks"]),
        ("Completed", report_data["summary"]["completed"]),
        ("In Progress", report_data["summary"]["in_progress"]),
        ("TODO", report_data["summary"]["todo"]),
        ("On Hold", report_data["summary"]["on_hold"]),
        ("Cancelled", report_data["summary"]["cancelled"]),
        ("Overdue", report_data["summary"]["overdue"]),
        ("Total Meetings", report_data["summary"]["total_meetings"]),
    ]
    start_row = 5
    summary_ws.cell(row=start_row, column=1, value="Metric").font = header_font
    summary_ws.cell(row=start_row, column=1).fill = header_fill
    summary_ws.cell(row=start_row, column=2, value="Value").font = header_font
    summary_ws.cell(row=start_row, column=2).fill = header_fill
    for i, (label, value) in enumerate(summary_rows, start=start_row + 1):
        summary_ws.cell(row=i, column=1, value=label)
        summary_ws.cell(row=i, column=2, value=value)
    summary_ws.column_dimensions["A"].width = 22
    summary_ws.column_dimensions["B"].width = 15

    def build_sheet(name, headers, rows, row_builder):
        ws = wb.create_sheet(name)
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for r, item in enumerate(rows, start=2):
            for c, value in enumerate(row_builder(item), start=1):
                ws.cell(row=r, column=c, value=value)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        ws.freeze_panes = "A2"
        return ws

    build_sheet("Tasks", TASK_HEADERS, report_data["tasks"], _task_row)
    build_sheet("Meetings", MEETING_HEADERS, report_data["meetings"], _meeting_row)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ------------------------------------------------------------------ CSV ----

def generate_csv_zip(report_data):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        tasks_csv = io.StringIO()
        writer = csv.writer(tasks_csv)
        writer.writerow(TASK_HEADERS)
        for t in report_data["tasks"]:
            writer.writerow(_task_row(t))
        zf.writestr("tasks.csv", tasks_csv.getvalue())

        meetings_csv = io.StringIO()
        writer = csv.writer(meetings_csv)
        writer.writerow(MEETING_HEADERS)
        for m in report_data["meetings"]:
            writer.writerow(_meeting_row(m))
        zf.writestr("meetings.csv", meetings_csv.getvalue())

    buffer.seek(0)
    return buffer


# ------------------------------------------------------------------ PDF ----

def _add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(
        landscape(letter)[0] - 0.5 * inch, 0.4 * inch, f"Page {doc.page}"
    )
    canvas.restoreState()


def generate_pdf(report_data):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8, leading=10)
    title_style = styles["Title"]
    heading_style = styles["Heading2"]

    story = []
    story.append(Paragraph("Daily Task &amp; Meeting Tracker", title_style))
    story.append(Paragraph(f"Report Period: {report_data['range_label']}", styles["Normal"]))
    story.append(
        Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}", styles["Normal"])
    )
    story.append(Spacer(1, 12))

    summary = report_data["summary"]
    summary_table_data = [
        ["Total Tasks", "Completed", "In Progress", "TODO", "On Hold", "Cancelled", "Overdue", "Meetings"],
        [
            summary["total_tasks"],
            summary["completed"],
            summary["in_progress"],
            summary["todo"],
            summary["on_hold"],
            summary["cancelled"],
            summary["overdue"],
            summary["total_meetings"],
        ],
    ]
    summary_table = Table(summary_table_data, hAlign="LEFT")
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    story.append(summary_table)
    story.append(Spacer(1, 16))

    def escape(text):
        return Paragraph((text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), cell_style)

    story.append(Paragraph("Task Report", heading_style))
    if report_data["tasks"]:
        data = [TASK_HEADERS] + [
            [escape(str(v)) for v in _task_row(t)] for t in report_data["tasks"]
        ]
        table = Table(data, repeatRows=1, colWidths=[0.7 * inch, 0.5 * inch, 1.6 * inch, 0.7 * inch, 0.9 * inch, 0.7 * inch, 2.5 * inch])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("No tasks found for this date range.", styles["Normal"]))

    story.append(Spacer(1, 16))
    story.append(Paragraph("Meeting Report", heading_style))
    if report_data["meetings"]:
        data = [MEETING_HEADERS] + [
            [escape(str(v)) for v in _meeting_row(m)] for m in report_data["meetings"]
        ]
        table = Table(data, repeatRows=1, colWidths=[0.7 * inch, 0.5 * inch, 1.3 * inch, 1.4 * inch, 1.4 * inch, 1.4 * inch, 1.5 * inch])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("No meetings found for this date range.", styles["Normal"]))

    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    buffer.seek(0)
    return buffer
